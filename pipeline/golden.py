"""Golden-50 human labeling helper. No model calls.

    python -m pipeline.golden sheet    # data/golden_50_to_label.csv -> labels/golden_50_labeling.xlsx (dropdowns)
    python -m pipeline.golden export   # labeled xlsx -> evals/golden_50_human.csv, after validation

These human labels are evaluation answers only. They must never be sent to a model,
used as prompt examples, or used to tune thresholds.
"""

import csv
import sys

from pipeline.common import DATA, INTENTS, ROOT, TOPICS, csv_rows

SOURCE = DATA / "golden_50_to_label.csv"
SHEET = ROOT / "labels" / "golden_50_labeling.xlsx"
OUT = ROOT / "evals" / "golden_50_human.csv"

SOURCE_COLS = ["review_id", "review_text", "review_rating", "review_likes", "app_version", "review_timestamp"]
LABEL_COLS = ["topic", "intent", "sentiment", "severity", "entities", "evidence_quote", "needs_review", "notes"]
CHOICES = {
    "topic": TOPICS,
    "intent": ("cancellation", "complaint", "request", "praise", "unclear"),  # precedence order
    "sentiment": ("-1", "-0.5", "0", "0.5", "1"),
    "severity": ("1", "2", "3", "4", "5"),
    "needs_review": ("FALSE", "TRUE"),
}
RULES = [
    ("TOPIC", "Pick the problem with the highest severity; on a tie, the first one mentioned."),
    ("", "Praise: the first specific praised feature. General praise ('great app') = other."),
    ("", "A paid-plan mention alone is NOT billing. Premium failing to activate IS billing. Crash for a paying user = playback."),
    ("access", "Login, signup, password or account access"),
    ("usability", "Navigation, controls, layout, queue/playlist management, ad interruptions"),
    ("playback", "Playback failure, crashes, lag, connection failures, audio quality, resource use"),
    ("downloads", "Downloading, saved music, offline listening, disappearing downloads"),
    ("catalog", "Missing songs/artists, search/discovery, recommendations, lyrics availability"),
    ("billing", "Price, charges, subscriptions, paywalls, premium entitlement, premium-only controls"),
    ("support", "Contacting support and the support response"),
    ("other", "General praise/criticism, unrelated content, or no supported specific topic"),
    ("", ""),
    ("INTENT", "Use the FIRST that applies: cancellation > complaint > request > praise > unclear"),
    ("cancellation", "Explicitly leaving, uninstalling, cancelling, or threatening to"),
    ("complaint", "Negative experience, including mixed praise + criticism. 'Bad app' is a complaint."),
    ("request", "Wants a change, no reported failure"),
    ("praise", "Positive, no complaint"),
    ("unclear", "Meaningless/unrelated text, bare boycott slogans (unless a product complaint or personal departure)"),
    ("", ""),
    ("SEVERITY", "Stars, angry words or 'I'm cancelling' do NOT raise severity by themselves."),
    ("1", "No reported problem: praise, neutral/unclear, or a pure feature request"),
    ("2", "Dislike, generic criticism, minor annoyance, cosmetic; no functional loss"),
    ("3", "Degraded/restricted function; some use or workaround remains"),
    ("4", "A core task clearly blocked (can't log in, can't play music)"),
    ("5", "Explicit serious financial, privacy or data harm (e.g. charged after cancelling)"),
    ("", ""),
    ("SENTIMENT", "-1 very negative ... 0 neutral/mixed ... 1 very positive. From the text, not the stars."),
    ("EVIDENCE", "Copy an EXACT piece of the review text that supports your label (short review = whole text)."),
    ("ENTITIES", "Optional: features/things named, separated by ; (e.g. 'shuffle; lyrics'). Leave blank if none."),
    ("NEEDS_REVIEW", "TRUE if the review is genuinely ambiguous or you lack context. Explain in notes."),
]


def make_sheet():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = list(csv_rows(SOURCE))
    wb = Workbook()
    ws = wb.active
    ws.title = "label"
    headers = ["#"] + SOURCE_COLS + LABEL_COLS
    ws.append(headers)
    for i, r in enumerate(rows, 1):
        ws.append([i] + [r[c] for c in SOURCE_COLS] + ["" for _ in LABEL_COLS])

    fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    widths = {"#": 4, "review_id": 12, "review_text": 70, "review_rating": 7, "review_likes": 7, "app_version": 10,
              "review_timestamp": 11, "topic": 12, "intent": 13, "sentiment": 10, "severity": 9, "entities": 18,
              "evidence_quote": 40, "needs_review": 13, "notes": 30}
    for idx, name in enumerate(headers, 1):
        col = ws.cell(1, idx).column_letter
        ws.column_dimensions[col].width = widths[name]
        if name in LABEL_COLS:
            ws.cell(1, idx).fill = fill
        if name in CHOICES:
            dv = DataValidation(type="list", formula1='"' + ",".join(CHOICES[name]) + '"', allow_blank=True,
                                showErrorMessage=True, errorTitle="Invalid", error="Pick from the list")
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{len(rows) + 1}")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"

    rules = wb.create_sheet("rules")
    for a, b in RULES:
        rules.append([a, b])
    rules.column_dimensions["A"].width = 15
    rules.column_dimensions["B"].width = 110
    for row in rules.iter_rows():
        if row[0].value and row[0].value.isupper():
            row[0].font = row[1].font = Font(bold=True)

    SHEET.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SHEET)
    print(f"Wrote {SHEET.relative_to(ROOT)} ({len(rows)} reviews). Fill the yellow columns, save, then run export.")


def export():
    from openpyxl import load_workbook

    source = {r["review_id"]: r for r in csv_rows(SOURCE)}
    ws = load_workbook(SHEET, data_only=True)["label"]
    header = [c.value for c in ws[1]]
    problems, out = [], []
    for n, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        row = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, values)}
        rid = row["review_id"]
        if not rid:
            continue
        src = source.get(rid)
        if src is None:
            problems.append(f"row {n}: unknown review_id {rid}")
            continue
        for col in ("topic", "intent", "severity", "sentiment"):
            if not row[col]:
                problems.append(f"row {n}: {col} is blank")
        for col, allowed in CHOICES.items():
            if row[col] and row[col] not in allowed and not (col == "needs_review" and row[col].upper() in allowed):
                problems.append(f"row {n}: {col}={row[col]!r} not allowed")
        if row["sentiment"]:
            try:
                assert -1 <= float(row["sentiment"]) <= 1
            except (ValueError, AssertionError):
                problems.append(f"row {n}: sentiment must be a number from -1 to 1")
        quote = row["evidence_quote"]
        if not quote:
            problems.append(f"row {n}: evidence_quote is blank")
        elif quote not in src["review_text"]:  # Excel strips nothing inside the cell, but check exactness
            problems.append(f"row {n}: evidence_quote is not an exact substring of the review text")
        out.append({**{c: src[c] for c in SOURCE_COLS}, **{c: row[c] for c in LABEL_COLS},
                    "needs_review": (row["needs_review"] or "FALSE").lower()})

    if len(out) != len(source):
        problems.append(f"expected {len(source)} labeled rows, found {len(out)}")
    if problems:
        print("Not exported - fix these first:")
        print("\n".join("  " + p for p in problems))
        sys.exit(1)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=SOURCE_COLS + LABEL_COLS, lineterminator="\n")
        w.writeheader()
        w.writerows(out)
    print(f"All {len(out)} labels valid. Wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    {"sheet": make_sheet, "export": export}.get(sys.argv[1] if len(sys.argv) > 1 else "", lambda: print(__doc__))()
